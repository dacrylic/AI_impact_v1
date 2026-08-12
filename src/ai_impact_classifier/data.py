from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Tuple

import pandas as pd
from openpyxl import load_workbook


TEXT_COLUMNS = ("jobrole_title", "keytask_content")
GROUP_COLUMN = "jobrole_id"
RAW_LABEL_COLUMN = "openai_label"
MODEL_LABEL_COLUMN = "model_label"
# LABEL_COLUMN is the only label column modeling code should use.
LABEL_COLUMN = MODEL_LABEL_COLUMN
SCORE_COLUMN = "ai_impact_score"
MODEL_LABEL_ORDER = ["E0", "E1", "E23"]
RAW_TO_MODEL_LABEL = {"E0": "E0", "E1": "E1", "E2": "E23", "E3": "E23"}


@dataclass(frozen=True)
class DatasetBundle:
    frame: pd.DataFrame
    label_order: List[str]
    label_to_rank: Dict[str, int]
    rank_to_label: Dict[int, str]
    score_by_label: Dict[str, float]


def load_workbook_frame(path: str | Path) -> pd.DataFrame:
    path = Path(path)
    frame = pd.read_excel(path, sheet_name="task ai impact", engine="openpyxl")
    frame = frame.copy()
    frame = frame.dropna(subset=[RAW_LABEL_COLUMN]).reset_index(drop=True)
    frame[MODEL_LABEL_COLUMN] = frame[RAW_LABEL_COLUMN].map(RAW_TO_MODEL_LABEL)
    if frame[MODEL_LABEL_COLUMN].isna().any():
        unknown = frame.loc[frame[MODEL_LABEL_COLUMN].isna(), RAW_LABEL_COLUMN].dropna().unique().tolist()
        raise ValueError(f"Unsupported raw labels for the canonical model target: {unknown}")
    frame[TEXT_COLUMNS[0]] = frame[TEXT_COLUMNS[0]].fillna("").astype(str)
    frame[TEXT_COLUMNS[1]] = frame[TEXT_COLUMNS[1]].fillna("").astype(str)
    return frame


def infer_label_order(frame: pd.DataFrame) -> Tuple[List[str], Dict[str, float]]:
    # E23 deliberately has no single pseudo-continuous score. The source score remains
    # available as an audit column, but it is not a modeling target or feature.
    return MODEL_LABEL_ORDER.copy(), {}


def build_bundle(path: str | Path) -> DatasetBundle:
    frame = load_workbook_frame(path)
    label_order, score_by_label = infer_label_order(frame)
    label_to_rank = {label: idx for idx, label in enumerate(label_order)}
    rank_to_label = {idx: label for label, idx in label_to_rank.items()}
    frame["label_rank"] = frame[LABEL_COLUMN].map(label_to_rank).astype(int)
    frame["target_text"] = (
        frame["jobrole_title"].str.strip() + "\n" + frame["keytask_content"].str.strip()
    ).str.strip()
    return DatasetBundle(
        frame=frame,
        label_order=label_order,
        label_to_rank=label_to_rank,
        rank_to_label=rank_to_label,
        score_by_label=score_by_label,
    )
